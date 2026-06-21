from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path("KLASIFIKASI KINERJA KARYAWAN TAHUNAN MENGGUNAKAN METODE 1.docx")
OUTPUT = Path(
    "KLASIFIKASI KINERJA KARYAWAN TAHUNAN MENGGUNAKAN METODE 1 - revisi abstrak.docx"
)

ABSTRAK_META = (
    'Markelius Hendra Manalu, 2055201076. "Klasifikasi Kinerja Karyawan Tahunan '
    'Menggunakan Metode K-Nearest Neighbor pada PT. Zetwell Quantum Solution." '
    "Di bawah bimbingan Ari Sellyana, M.Kom dan Ir. Mizztazihim Suhaidi, M.Kom, "
    "Mei 2026, 120 halaman + XVI. "
)
ABSTRAK_BODY = (
    "Penilaian kinerja karyawan di PT. Zetwell Quantum Solution masih menghadapi "
    "kendala pada objektivitas klasifikasi dan pengelolaan data yang belum "
    "terintegrasi. Penelitian ini bertujuan membangun sistem berbasis web yang "
    "menerapkan metode K-Nearest Neighbor (KNN) untuk membantu klasifikasi "
    "kinerja karyawan tahunan secara lebih terstruktur, objektif, dan efisien. "
    "Metode pengembangan sistem yang digunakan adalah waterfall, sedangkan proses "
    "klasifikasi dilakukan dengan KNN menggunakan Euclidean Distance. Data latih "
    "yang digunakan berjumlah 80 data hasil penilaian tervalidasi periode 2021 "
    "sampai 2025 dengan 13 kriteria penilaian dan parameter k = 3. Hasil "
    "penelitian menunjukkan bahwa sistem yang dibangun mampu mengelola data "
    "karyawan, kriteria penilaian, relasi atasan, periode penilaian, proses "
    "penilaian, analisa klasifikasi, dan laporan dalam satu platform terintegrasi. "
    "Pada pengujian data uji, hasil voting mayoritas dari tiga tetangga terdekat "
    "menghasilkan klasifikasi kinerja pada kategori Baik. Dengan demikian, "
    "penerapan metode KNN pada sistem ini dapat mendukung proses penilaian "
    "kinerja karyawan yang lebih cepat, terdokumentasi, dan mudah ditelusuri kembali."
)
ABSTRAK_KEYWORDS = (
    "Kata Kunci : K-Nearest Neighbor, Klasifikasi Kinerja, Penilaian Karyawan, "
    "Sistem Berbasis Web"
)

ABSTRACT_META = (
    'Markelius Hendra Manalu, 2055201076. "Annual Employee Performance '
    'Classification Using the K-Nearest Neighbor Method at PT. Zetwell Quantum '
    'Solution." Under the guidance of Ari Sellyana, M.Kom and Ir. Mizztazihim '
    "Suhaidi, M.Kom, May 2026, 120 pages + XVI. "
)
ABSTRACT_BODY = (
    "Employee performance assessment at PT. Zetwell Quantum Solution still faces "
    "challenges in classification objectivity and non-integrated data management. "
    "This study aims to build a web-based system that applies the K-Nearest "
    "Neighbor (KNN) method to support annual employee performance classification "
    "in a more structured, objective, and efficient manner. The system was "
    "developed using the waterfall model, while the classification process "
    "employed KNN with Euclidean Distance. The training data consisted of 80 "
    "validated performance records from 2021 to 2025, 13 assessment criteria, "
    "and k = 3. The results show that the developed system can manage employee "
    "data, assessment criteria, supervisor relationships, assessment periods, "
    "appraisal processes, classification analysis, and reporting in one integrated "
    "platform. In the test data evaluation, majority voting from the three nearest "
    "neighbors produced a performance classification in the Good category. "
    "Therefore, the implementation of KNN in this system can support a faster, "
    "well-documented, and more traceable employee performance assessment process."
)
ABSTRACT_KEYWORDS = (
    "Keywords : K-Nearest Neighbor, Employee Performance Classification, "
    "Performance Assessment, Web-Based System"
)


def set_run_font(run, *, bold=None, italic=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(key), "Times New Roman")


def format_paragraph(paragraph, *, alignment, line_spacing=1.0, first_line_indent=0):
    paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(0)
    fmt.right_indent = Inches(0)
    fmt.first_line_indent = Inches(first_line_indent)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = line_spacing


def insert_page_break(anchor):
    paragraph = anchor.insert_paragraph_before()
    format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    run = paragraph.add_run()
    set_run_font(run)
    run.add_break(WD_BREAK.PAGE)


def insert_heading(anchor, text):
    paragraph = anchor.insert_paragraph_before()
    paragraph.style = "Style1"
    format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph.paragraph_format.space_after = Pt(12)
    run = paragraph.add_run(text)
    set_run_font(run, bold=True)


def insert_body(anchor, meta_text, body_text, *, meta_bold=False, meta_italic=False, body_italic=False):
    paragraph = anchor.insert_paragraph_before()
    paragraph.style = "Normal"
    format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    meta_run = paragraph.add_run(meta_text)
    set_run_font(meta_run, bold=meta_bold, italic=meta_italic)
    body_run = paragraph.add_run(body_text)
    set_run_font(body_run, italic=body_italic)


def insert_keywords(anchor, text, *, italic=False):
    paragraph = anchor.insert_paragraph_before()
    paragraph.style = "Normal"
    format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    paragraph.paragraph_format.space_before = Pt(12)
    run = paragraph.add_run(text)
    set_run_font(run, bold=True, italic=italic)


def is_section_break_paragraph(paragraph):
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.sectPr is not None


def find_front_matter_section_break(document):
    for index, paragraph in enumerate(document.paragraphs[:-1]):
        text = " ".join(paragraph.text.split())
        if text.startswith("BAB I"):
            candidate = document.paragraphs[index - 1]
            if is_section_break_paragraph(candidate):
                return candidate
    raise RuntimeError("Target section break before BAB I was not found.")


def main():
    shutil.copyfile(SOURCE, OUTPUT)
    document = Document(str(OUTPUT))
    anchor = find_front_matter_section_break(document)

    insert_page_break(anchor)
    insert_heading(anchor, "ABSTRAK")
    insert_body(
        anchor,
        ABSTRAK_META,
        ABSTRAK_BODY,
        meta_bold=True,
        meta_italic=False,
        body_italic=False,
    )
    insert_keywords(anchor, ABSTRAK_KEYWORDS, italic=False)

    insert_page_break(anchor)
    insert_heading(anchor, "ABSTRACT")
    insert_body(
        anchor,
        ABSTRACT_META,
        ABSTRACT_BODY,
        meta_bold=True,
        meta_italic=True,
        body_italic=True,
    )
    insert_keywords(anchor, ABSTRACT_KEYWORDS, italic=True)

    document.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
