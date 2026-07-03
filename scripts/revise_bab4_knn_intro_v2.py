from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SOURCE = Path("KLASIFIKASI KINERJA KARYAWAN TAHUNAN MENGGUNAKAN METODE 1 - revisi abstrak.docx")
OUTPUT = Path(
    "KLASIFIKASI KINERJA KARYAWAN TAHUNAN MENGGUNAKAN METODE 1 - revisi bab4 knn.docx"
)


PARAGRAPHS = [
    [
        ("Secara konseptual, ", False),
        ("K-Nearest Neighbor", True),
        (" merupakan algoritma klasifikasi ", False),
        ("supervised", True),
        (" yang bekerja dengan membandingkan data uji terhadap kumpulan ", False),
        ("data latih", True),
        (
            ". Setiap data uji akan dihitung tingkat kedekatannya dengan seluruh data latih, "
            "kemudian dipilih sejumlah k tetangga terdekat sebagai dasar penentuan kelas.",
            False,
        ),
    ],
    [
        ("Dalam penelitian ini, ukuran kedekatan antar data dihitung menggunakan ", False),
        ("Euclidean Distance", True),
        (
            ". Semakin kecil nilai jarak yang diperoleh, maka semakin tinggi tingkat kemiripan "
            "antara data uji dan data latih. Kelas yang paling dominan dari tetangga terdekat "
            "tersebut selanjutnya ditetapkan sebagai hasil klasifikasi kinerja karyawan.",
            False,
        ),
    ],
]


def insert_paragraph_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def clone_paragraph_style(source_paragraph, target_paragraph):
    target_paragraph.style = source_paragraph.style
    target_paragraph.alignment = source_paragraph.alignment
    target_format = target_paragraph.paragraph_format
    source_format = source_paragraph.paragraph_format
    target_format.left_indent = source_format.left_indent
    target_format.right_indent = source_format.right_indent
    target_format.first_line_indent = source_format.first_line_indent
    target_format.space_before = source_format.space_before
    target_format.space_after = source_format.space_after
    target_format.line_spacing = source_format.line_spacing


def set_run_font(run, italic=False):
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(key), "Times New Roman")


def find_anchor(document):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(
            "Hasil penilaian kinerja yang diberikan oleh pengawas (supervisor)"
        ):
            return paragraph
    raise RuntimeError("Anchor paragraph in Bab IV not found.")


def append_after(paragraph, content_chunks):
    new_paragraph = insert_paragraph_after(paragraph)
    clone_paragraph_style(paragraph, new_paragraph)
    for text, italic in content_chunks:
        run = new_paragraph.add_run(text)
        set_run_font(run, italic=italic)
    return new_paragraph


def main():
    shutil.copyfile(SOURCE, OUTPUT)
    document = Document(str(OUTPUT))
    anchor = find_anchor(document)

    current = anchor
    for chunks in PARAGRAPHS:
        current = append_after(current, chunks)

    document.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
